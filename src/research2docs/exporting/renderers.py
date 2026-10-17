"""No model calls, raw HTML, active links, embedded files or executable OOXML."""

import html
import io
import json
import re
import zipfile
from datetime import UTC
from typing import Any

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from jinja2 import Environment, StrictUndefined

MAX_EXPORT_BYTES = 32 * 1024 * 1024


def display(text: str) -> str:
    return "".join(
        c
        if c in "\t\n\r"
        or 0x20 <= ord(c) <= 0xD7FF
        or 0xE000 <= ord(c) <= 0xFFFD
        or 0x10000 <= ord(c) <= 0x10FFFF
        else f"[U+{ord(c):04X}]"
        for c in text
    )


def blocks(manifest: dict[str, Any]) -> list[tuple[str, str]]:
    draft = manifest["draft"]
    ledger = manifest["ledger"]
    units = ledger["units"]
    unresolved = sum(u["claim"]["human_review"] != "accepted" for u in units)
    result = [
        ("title", draft["title"]),
        ("p", "DRAFT — source references and review status preserved"),
        (
            "p",
            f"Revision {draft['revision_number']} · {draft['origin']} · Revision ID {draft['revision_id']}",
        ),
        (
            "p",
            f"{unresolved} of {len(units)} claim units remain unresolved. Human decisions are self-reported editorial judgments, not independent proof of truth.",
        ),
        (
            "p",
            "AI support is provisional. Source relevance and citation existence do not establish factual correctness.",
        ),
        ("h", "Retained warnings"),
    ]
    result += [("p", warning) for warning in draft["warnings"]]
    result += [
        (
            "p",
            "Readable exports display XML-invalid control characters as U+ labels; the JSON manifest preserves original strings.",
        )
    ]
    indexed = {(u["section_index"], u["block_index"]): u for u in units}
    for si, section in enumerate(draft["content"]["sections"]):
        result.append(("h", section["heading"]))
        for bi, block in enumerate(section["blocks"]):
            unit = indexed[(si, bi)]
            claim = unit["claim"]
            result += [
                ("p", block["text"]),
                (
                    "p",
                    f"Review: {claim['assessment']} · Human: {claim['human_review']} · Claim {claim['claim_id']}",
                ),
            ]
            result.append(
                (
                    "p",
                    "References: "
                    + (", ".join(block["citations"]) or "NONE — requires verification"),
                )
            )
            if block["review_note"]:
                result.append(("p", "Review note: " + block["review_note"]))
    result.append(("h", "Document limitations"))
    result += [("p", note) for note in draft["content"]["limitations"]]
    result.append(("h", "Claim ledger"))
    for i, unit in enumerate(units):
        claim = unit["claim"]
        result += [
            ("h", f"Claim {i + 1} {claim['claim_id']}"),
            ("p", claim["text"]),
            (
                "p",
                f"Assessment: {claim['assessment']} · State: {unit['state']} · Human review: {claim['human_review']}",
            ),
            ("p", "Assessment explanation: " + claim["explanation"]),
            (
                "p",
                "Human rationale: "
                + (claim["human_rationale"] or "No current human rationale recorded."),
            ),
            (
                "p",
                f"Automated method/time retained: {claim['assessment_method']} · {claim['assessed_at']}",
            ),
        ]
        result += [("p", flag) for flag in unit["review_flags"]]
        for kind in ["supporting", "contradictory", "temporal"]:
            for quote in unit[kind]:
                result += [
                    (
                        "p",
                        f"{kind.title()} · Passage {quote['passage_id']} · Characters [{quote['start']}, {quote['end']})",
                    ),
                    ("pre", quote["quote"]),
                ]
        for pid in unit.get("human_evidence", []):
            result.append(("p", "Human-reviewed passage: " + pid))
    result.append(("h", "Evidence passages and source metadata"))
    aliases = {e["passage"]["passage_id"]: e["alias"] for e in draft["evidence"]}
    for item in manifest["evidence"]:
        passage = item["passage"]
        source = item["source"]
        pid = passage["passage_id"]
        result += [
            ("h", (aliases.get(pid, "Additional reviewed evidence") + " · " + source["title"])),
            ("p", f"Passage: {pid} · Source: {source['source_id']} · File: {source['filename']}"),
            (
                "p",
                f"Author: {source['author'] or 'Unavailable'} · Imported: {source['imported_at']} · SHA-256: {source['sha256']}",
            ),
            (
                "p",
                f"Page: {passage['page'] or 'Unavailable'} · Section: {passage['section'] or 'Unavailable'} · Extraction: {passage['extraction_method']}",
            ),
            ("p", "Reference URL label only: " + (source["supplied_url"] or "None")),
            ("pre", passage["text"]),
        ]
    result.append(("h", "Editorial decision history"))
    if not manifest["human_decisions"]:
        result.append(("p", "No human decisions recorded for this revision."))
    for event in manifest["human_decisions"]:
        review = event["payload"]["review"]
        result.append(
            (
                "p",
                f"{event['created_at']} · {event['action']} · {review['reviewer']} · {review['rationale']}",
            )
        )
    result += [
        ("h", "Generation provenance"),
        (
            "p",
            f"Model: {draft['model']} · Digest: {draft['model_digest']} · Prompt: {draft['prompt_version']} · Prompt SHA-256: {draft['prompt_sha256']}",
        ),
        (
            "pre",
            "Generation request parameters: "
            + json.dumps(manifest["generation_run"]["parameters"], sort_keys=True),
        ),
        (
            "p",
            "The JSON evidence manifest contains the exact revision, complete claim records and stored assessment/decision history. Export does not change any review status.",
        ),
    ]
    return [(kind, display(text)) for kind, text in result]


def markdown(parts: list[tuple[str, str]]) -> bytes:
    def escape(text: str) -> str:
        text = html.escape(text, quote=False)
        return re.sub(r"([\\`*_{}\[\]()#+.!|>\-])", r"\\\1", text)

    out = []
    for kind, text in parts:
        if kind == "pre":
            ticks = "`" * max(
                3, max((len(m.group()) for m in re.finditer(r"`+", text)), default=0) + 1
            )
            out.append(ticks + "text\n" + text + "\n" + ticks)
        else:
            out.append(("# " if kind == "title" else "## " if kind == "h" else "") + escape(text))
    return ("\n\n".join(out) + "\n").encode("utf-8")


HTML = """<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><meta http-equiv="Content-Security-Policy" content="default-src 'none'; style-src 'unsafe-inline'; base-uri 'none'; form-action 'none'"><title>{{ title }}</title><style>body{font:16px/1.55 system-ui,sans-serif;max-width:900px;margin:40px auto;padding:0 24px;color:#17202a}h1{font-size:30px}h2{font-size:21px;margin-top:28px}p,pre{white-space:pre-wrap;overflow-wrap:anywhere}pre{font:14px/1.5 monospace;background:#f3f4f6;padding:12px}h1,h2{color:#000}p{margin:10px 0}</style></head><body>{% for kind,text in parts %}{% if kind=='title' %}<h1>{{ text }}</h1>{% elif kind=='h' %}<h2>{{ text }}</h2>{% elif kind=='pre' %}<pre>{{ text }}</pre>{% else %}<p>{{ text }}</p>{% endif %}{% endfor %}</body></html>"""


def html_document(parts: list[tuple[str, str]]) -> bytes:
    env = Environment(autoescape=True, undefined=StrictUndefined)
    return env.from_string(HTML).render(title=parts[0][1], parts=parts).encode("utf-8")


def docx_document(parts: list[tuple[str, str]], created: Any) -> bytes:
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(0.85)
    for name in ["Normal", "Title", "Heading 1", "Heading 2"]:
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(6)
    for border in document.styles["Title"].element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    document.styles["Normal"].font.size = Pt(10)
    document.styles["Normal"].paragraph_format.line_spacing = 1.08
    document.styles["Title"].font.size = Pt(24)
    document.styles["Heading 1"].font.size = Pt(15)
    document.styles["Heading 2"].font.size = Pt(12)
    for kind, text in parts:
        if kind == "title":
            document.add_paragraph(text, "Title")
        elif kind == "h":
            document.add_paragraph(text, "Heading 1")
        else:
            paragraph = document.add_paragraph(text)
            paragraph.paragraph_format.widow_control = True
            if kind == "pre":
                paragraph.paragraph_format.left_indent = Inches(0.15)
                for run in paragraph.runs:
                    run.font.name = "Consolas"
                    run.font.size = Pt(9)
    properties = document.core_properties
    properties.title = parts[0][1][:255]
    properties.author = "Research2Docs AI"
    properties.last_modified_by = "Research2Docs AI"
    properties.subject = "Draft with evidence and review status"
    properties.comments = ""
    stamp = created.astimezone(UTC).replace(tzinfo=None)
    properties.created = stamp
    properties.modified = stamp
    buffer = io.BytesIO()
    document.save(buffer)
    # Canonical ZIP timestamps/member ordering make repeated exports byte-reproducible.
    result = io.BytesIO()
    with (
        zipfile.ZipFile(buffer) as incoming,
        zipfile.ZipFile(result, "w", compression=zipfile.ZIP_DEFLATED) as outgoing,
    ):
        for name in sorted(incoming.namelist()):
            info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            outgoing.writestr(info, incoming.read(name))
    return result.getvalue()


def json_manifest(manifest: dict[str, Any]) -> bytes:
    return (json.dumps(manifest, ensure_ascii=True, sort_keys=True, indent=2) + "\n").encode(
        "utf-8"
    )
