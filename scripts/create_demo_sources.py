"""Create original fictional five-format source fixtures once; never overwrite files."""

from datetime import UTC, datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]


def main() -> None:
    import pymupdf
    from docx import Document

    folder = ROOT / "samples/fictional_research"
    folder.mkdir(parents=True, exist_ok=True)
    names = ["overview.md", "notes.txt", "reference.html", "workflow.docx", "troubleshooting.pdf"]
    if any((folder / name).exists() for name in names):
        raise SystemExit("Demo source files already exist; no files were overwritten.")
    (folder / "overview.md").write_text(
        "# Orion Archive — fictional research overview\n\n"
        "## Evidence marker\nRESEARCH_READY identifies the fictional documentation review stage.\n\n"
        "## Scope\nOrion Archive is invented demonstration material, not a real system. "
        "This overview does not establish deployment requirements or performance results.\n"
    )
    (folder / "notes.txt").write_text(
        "Fictional research notes\n\nRESEARCH_READY is a documentation label, not a security certification.\n"
        "Retention periods and recovery procedures are intentionally unspecified.\n"
    )
    (folder / "reference.html").write_text(
        "<!doctype html><html><head><title>Orion Archive reference — fictional</title>"
        '<meta name="author" content="Fictional Research Team"></head><body>'
        "<h1>Evidence states</h1><p>RESEARCH_READY means the fictional source set is available for review.</p>"
        "<h2>Boundaries</h2><p>Source relevance does not prove factual correctness.</p>"
        "<script>/* Inert fixture: this script must not execute or become evidence. */</script>"
        "</body></html>"
    )
    doc = Document()
    doc.core_properties.title = "Orion Archive workflow — fictional"
    doc.core_properties.author = "Fictional Research Team"
    doc.core_properties.created = datetime(2026, 10, 17, tzinfo=UTC)
    doc.core_properties.modified = datetime(2026, 10, 17, tzinfo=UTC)
    doc.add_heading("Research review workflow", level=1)
    doc.add_paragraph("RESEARCH_READY labels a source collection awaiting human review.")
    doc.add_heading("Evidence handling", level=2)
    doc.add_paragraph(
        "Inspect the original context before relying on an excerpt. No automated verification result is claimed."
    )
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Marker"
    table.cell(0, 1).text = "RESEARCH_READY"
    doc.save(str(folder / "workflow.docx"))
    pdf = pymupdf.open()
    pdf.set_metadata(
        {"title": "Orion Archive troubleshooting — fictional", "author": "Fictional Research Team"}
    )
    pdf.new_page().insert_text(
        (72, 72),
        "Orion Archive: fictional troubleshooting reference\nDemonstration source only; no real product or deployment.",
    )
    pdf.new_page().insert_text(
        (72, 72),
        "Evidence marker\nRESEARCH_READY identifies source availability, not verified claims.\nRecovery and retention specifications are not established here.",
    )
    pdf.save(str(folder / "troubleshooting.pdf"))
    pdf.close()
    print("Created five original fictional source files. Existing originals were not modified.")


if __name__ == "__main__":
    main()
