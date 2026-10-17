"""Export integrity, inactive markup and complete saved review provenance."""

import io
import json
import zipfile

import pytest
from bs4 import BeautifulSoup
from docx import Document
from test_editorial import decide, edited
from test_verification import assess, first
from test_verification import review_fixture as _shared_review_fixture

from research2docs.core.errors import ExportError, GenerationError


@pytest.fixture(name="review_fixture")
def export_fixture(tmp_path, monkeypatch):
    return _shared_review_fixture.__wrapped__(tmp_path, monkeypatch)


@pytest.mark.parametrize("format", ["markdown", "docx", "html", "json"])
def test_all_formats_preserve_revision_citations_and_unresolved_flags(review_fixture, format):
    w, p, d, _ = review_fixture
    before = first(w, p, d)
    output = w.exports.export(p.project_id, d.draft_id, format)
    assert output.data and output.revision_id == d.revision_id
    assert "/" not in output.filename and "\\" not in output.filename
    assert str(d.draft_id) in output.filename and "-r1-" in output.filename
    if format == "docx":
        document = Document(io.BytesIO(output.data))
        assert not document.styles["Title"].element.xpath(".//w:pBdr")
        text = "\n".join(p.text for p in document.paragraphs)
        with zipfile.ZipFile(io.BytesIO(output.data)) as z:
            assert not any("vbaProject" in name or "embeddings/" in name for name in z.namelist())
            for name in z.namelist():
                if name.endswith(".rels"):
                    assert b'TargetMode="External"' not in z.read(name)
    elif format == "html":
        text = BeautifulSoup(output.data, "html.parser").get_text()
    elif format == "markdown":
        from markdown_it import MarkdownIt

        text = BeautifulSoup(MarkdownIt().render(output.data.decode()), "html.parser").get_text()
    else:
        text = output.data.decode()
    assert "pending" in text and "not_assessed" in text
    assert str(before.claim.claim_id) in text
    assert str(before.cited_passages[0]) in text
    assert "false" in text and "true" in text
    assert first(w, p, d) == before
    assert output.claim_versions[str(before.claim.claim_id)] == 0


@pytest.mark.parametrize("format", ["markdown", "docx", "html", "json"])
def test_repeated_export_is_byte_reproducible(review_fixture, format):
    w, p, d, _ = review_fixture
    assert (
        w.exports.export(p.project_id, d.draft_id, format).data
        == w.exports.export(p.project_id, d.draft_id, format).data
    )


def test_manifest_preserves_human_decisions_automatic_history_and_exact_sources(review_fixture):
    w, p, d, _ = review_fixture
    assessed = assess(w, p, d)
    accepted = decide(w, p, d, assessed)
    data = json.loads(w.exports.export(p.project_id, d.draft_id, "json").data)
    assert data["export_status"] == "draft"
    assert data["generation_run"]["parameters"]["context_tokens"] == 32768
    assert data["ledger"]["units"][0]["claim"]["human_review"] == "accepted"
    assert data["ledger"]["units"][0]["automated_assessment"] == "source_supported"
    assert data["human_decisions"][0]["payload"]["review"]["reviewer"] == "Fictional Test Reviewer"
    assert data["assessment_history"][0]["status"] == "completed"
    evidence = {e["passage"]["passage_id"]: e for e in data["evidence"]}
    for quote in accepted.supporting:
        assert (
            evidence[str(quote.passage_id)]["passage"]["text"][quote.start : quote.end]
            == quote.quote
        )
    assert all(
        e["source"]["sha256"] and e["passage"]["project_id"] == str(p.project_id)
        for e in data["evidence"]
    )


def test_html_markdown_and_docx_keep_generated_markup_inert(review_fixture):
    w, p, d, _ = review_fixture
    attack = '![image](https://example.invalid/no-fetch.png) <script>alert(1)</script> <img src="https://example.invalid/x">'
    changed = w.editorial.save(p.project_id, d.draft_id, d.revision_id, edited(d, attack))
    html = w.exports.export(p.project_id, d.draft_id, "html").data
    soup = BeautifulSoup(html, "html.parser")
    assert not soup.find_all(["script", "img", "iframe", "a", "object", "form"])
    assert soup.find("meta", attrs={"http-equiv": "Content-Security-Policy"})
    assert attack in soup.get_text()
    md = w.exports.export(p.project_id, d.draft_id, "markdown").data.decode()
    assert "&lt;script&gt;" in md and "\\!\\[image\\]" in md
    document = Document(io.BytesIO(w.exports.export(p.project_id, d.draft_id, "docx").data))
    assert any(p.text == attack for p in document.paragraphs)
    assert changed.revision_number == 2


def test_historical_export_never_substitutes_latest_content(review_fixture):
    w, p, d, _ = review_fixture
    accepted = decide(w, p, d, first(w, p, d))
    latest = w.editorial.save(
        p.project_id, d.draft_id, d.revision_id, edited(d, "New revision text.")
    )
    historical = json.loads(
        w.exports.export(p.project_id, d.draft_id, "json", revision_id=d.revision_id).data
    )
    current = json.loads(w.exports.export(p.project_id, d.draft_id, "json").data)
    assert historical["draft"]["revision_id"] == str(d.revision_id)
    assert historical["ledger"]["units"][0]["claim"]["claim_id"] == str(accepted.claim.claim_id)
    assert historical["ledger"]["units"][0]["claim"]["human_review"] == "accepted"
    assert current["draft"]["revision_id"] == str(latest.revision_id)
    assert current["ledger"]["units"][0]["claim"]["human_review"] == "pending"
    assert not current["human_decisions"]


def test_source_urls_and_titles_are_labels_not_output_paths_or_links(review_fixture):
    w, p, d, _ = review_fixture
    title = "../../![escape](https://example.invalid/)"
    change = edited(d)
    change.title = title
    w.editorial.save(p.project_id, d.draft_id, d.revision_id, change)
    output = w.exports.export(p.project_id, d.draft_id, "html")
    assert ".." not in output.filename and "<h1>../../!" in output.data.decode()
    soup = BeautifulSoup(output.data, "html.parser")
    assert not soup.find_all(["a", "img"])


def test_audit_size_limit_fails_safely_instead_of_omitting_history(review_fixture):
    w, p, d, _ = review_fixture
    unit = first(w, p, d)
    with w.database.session(write=True) as c:
        c.execute(
            "INSERT INTO assessment_runs VALUES (?,?,?,?,?,?,?)",
            (
                str(p.project_id),
                "test-run",
                str(d.revision_id),
                str(unit.claim.claim_id),
                "failed",
                "fixture",
                json.dumps({"oversize": "x" * (8 * 1024 * 1024)}),
            ),
        )
    with pytest.raises(ExportError, match="history"):
        w.exports.export(p.project_id, d.draft_id, "json")


def test_export_selection_is_project_scoped_and_engine_has_no_network_dependency(
    review_fixture, monkeypatch
):
    import socket

    from research2docs.core.models import ProjectInput

    w, p, d, _ = review_fixture
    other = w.projects.create(ProjectInput(title="Other"))
    with pytest.raises(GenerationError):
        w.exports.export(other.project_id, d.draft_id, "json")
    monkeypatch.setattr(
        socket.socket, "connect", lambda *a, **k: pytest.fail("Export must not use network")
    )
    assert w.exports.export(p.project_id, d.draft_id, "html").data


def test_ui_prepares_and_invalidates_an_old_review_snapshot(review_fixture, monkeypatch):
    from pathlib import Path

    from streamlit.testing.v1 import AppTest

    w, p, d, _ = review_fixture
    monkeypatch.setenv("RESEARCH2DOCS_DATA_DIR", w.settings.data_dir)
    at = AppTest.from_file(
        str(Path(__file__).resolve().parents[2] / "ui/app.py"), default_timeout=30
    ).run()
    at.radio[0].set_value("Export Center").run()
    next(v for v in at.selectbox if v.label == "Export format").set_value(
        "JSON evidence manifest"
    ).run()
    next(b for b in at.button if b.label == "Prepare export").click().run()
    assert not at.exception and any("export prepared" in v.value for v in at.success)
    assert at.session_state["prepared_export"]["data"]
    decide(w, p, d, first(w, p, d))
    at.run()
    assert any("review status changed" in v.value for v in at.info)


def test_export_rejects_a_review_change_during_snapshot_preparation(review_fixture, monkeypatch):
    w, p, d, _ = review_fixture
    original = w.verification.ledger
    entered = False

    def raced(*args, **kwargs):
        nonlocal entered
        ledger = original(*args, **kwargs)
        if not entered:
            entered = True
            decide(w, p, d, ledger.units[0])
        return ledger

    monkeypatch.setattr(w.verification, "ledger", raced)
    with pytest.raises(ExportError, match="Review changed"):
        w.exports.export(p.project_id, d.draft_id, "json")


def test_control_characters_are_labeled_in_docx_and_exact_in_manifest(review_fixture):
    w, p, d, _ = review_fixture
    value = "Fictional text\x01 preserves the original code point."
    w.editorial.save(p.project_id, d.draft_id, d.revision_id, edited(d, value))
    manifest = json.loads(w.exports.export(p.project_id, d.draft_id, "json").data)
    assert manifest["draft"]["content"]["sections"][0]["blocks"][0]["text"] == value
    doc = Document(io.BytesIO(w.exports.export(p.project_id, d.draft_id, "docx").data))
    assert any("[U+0001]" in p.text for p in doc.paragraphs)


def test_full_controlled_workflow_import_review_edit_regenerate_export_reopen(
    review_fixture, five_sources, monkeypatch
):
    from test_editorial import regeneration_mock

    from research2docs.services.workspace_service import WorkspaceService

    w, p, d, _ = review_fixture
    for filename, data in five_sources.items():
        assert w.sources.import_document(p.project_id, data, filename).accepted
    assert len(w.sources.list_sources(p.project_id)) == 7
    assert w.sources.search(p.project_id, "EVIDENCE_READY")
    assessed = assess(w, p, d)
    decide(w, p, d, assessed)
    manual = w.editorial.save(p.project_id, d.draft_id, d.revision_id, edited(d))
    regeneration_mock(w, monkeypatch)
    regenerated = w.editorial.regenerate(
        p.project_id,
        d.draft_id,
        manual.revision_id,
        0,
        editor="Fictional Test Editor",
        note="Controlled complete-flow test.",
    )
    assert regenerated.revision_number == 3
    assert all(
        u.claim.human_review == "pending"
        for u in w.verification.ledger(p.project_id, d.draft_id).units
    )
    for kind in ["markdown", "docx", "html", "json"]:
        assert w.exports.export(p.project_id, d.draft_id, kind).data
    reopened = WorkspaceService(w.settings)
    assert len(reopened.editorial.history(p.project_id, d.draft_id)) == 3
    assert (
        reopened.verification.ledger(p.project_id, d.draft_id, revision_id=d.revision_id)
        .units[0]
        .claim.human_review
        == "accepted"
    )
    assert (
        json.loads(reopened.exports.export(p.project_id, d.draft_id, "json").data)["draft"][
            "revision_number"
        ]
        == 3
    )
