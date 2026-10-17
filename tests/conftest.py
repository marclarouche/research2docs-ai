"""Original fictional format fixtures, created locally; no downloads or real research."""

from io import BytesIO

import pytest


@pytest.fixture(scope="session")
def five_sources():
    import pymupdf
    from docx import Document

    document = Document()
    document.core_properties.title = "Fictional DOCX Research"
    document.core_properties.author = "Fictional Documentation Team"
    document.add_heading("Controlled workflow", level=1)
    document.add_paragraph("EVIDENCE_READY marks the fictional DOCX workflow.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Field"
    table.cell(0, 1).text = "Fictional value"
    binary = BytesIO()
    document.save(binary)
    pdf = pymupdf.open()
    pdf.new_page().insert_text((72, 72), "Fictional research overview")
    pdf.new_page().insert_text((72, 72), "EVIDENCE_READY marks the fictional PDF workflow.")
    pdf.set_metadata({"title": "Fictional PDF Research", "author": "Fictional Documentation Team"})
    pdf_bytes = pdf.tobytes()
    pdf.close()
    return {
        "reference.md": b"# Fictional Markdown Research\n\n## Evidence\nEVIDENCE_READY marks the fictional Markdown workflow.\n\n```text\n# This is code, not a heading\n```\n",
        "notes.txt": b"  EVIDENCE_READY marks the fictional plain text workflow.\r\nDo not infer real product behavior.\r\n",
        "reference.html": b'<!doctype html><html><head><title>Fictional HTML Research</title><meta name="author" content="Fictional Documentation Team"></head><body><h1>Evidence</h1><p>EVIDENCE_READY marks the fictional HTML workflow.</p><script>SHOULD_NOT_BE_EXTRACTED</script><style>secret-style</style><img src="https://example.invalid/no-fetch.png"><a href="https://example.invalid/">Reference label only</a></body></html>',
        "reference.docx": binary.getvalue(),
        "reference.pdf": pdf_bytes,
    }
